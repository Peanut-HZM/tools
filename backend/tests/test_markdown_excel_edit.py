"""
markdown_file_service.save_excel_edit 单元测试

覆盖 Excel 编辑保存接口的服务层行为：
  - set_cell 各 kind（auto/str/number/bool/formula）写入
  - insert_row / delete_row / insert_col / delete_col
  - add_sheet / rename_sheet / remove_sheet
  - 编辑保存后内嵌图片保留（真实场景：表格中带截图附件）
  - 原文件损坏防护（非法路径/非 xlsx/空操作/不存在文件）
  - 路径安全（.. 逃逸拒绝）
"""
import os
import sys
import shutil
import tempfile
import zipfile
from pathlib import Path

import pytest

# 支持从 backend/ 目录或 tests/ 目录运行
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import Workbook, load_workbook

from app.services.markdown_file_service import MarkdownFileService
from app.models.file_models import ExcelOperation


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------
@pytest.fixture
def temp_dir():
    """创建临时目录用于测试"""
    temp = tempfile.mkdtemp()
    yield temp
    shutil.rmtree(temp, ignore_errors=True)


@pytest.fixture
def file_service(temp_dir):
    """创建 MarkdownFileService 实例，root 指向 temp_dir"""
    service = MarkdownFileService("test-user-excel")
    service._root_path = Path(temp_dir).resolve()
    return service


def _make_xlsx(path: str, images: int = 0) -> None:
    """创建测试用 xlsx：Sheet1 两行数据 + 可选内嵌图片"""
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "Name"
    ws["B1"] = "Age"
    ws["A2"] = "Alice"
    ws["B2"] = 25
    for i in range(images):
        # 1x1 红色 PNG（最小合法 PNG 由 openpyxl 直接嵌入）
        from openpyxl.drawing.image import Image as XLImage
        import io
        # 生成最小 PNG：使用 openpyxl 自带的测试图片不可靠，手写 1x1 PNG 字节
        png = bytes.fromhex(
            "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
            "0000000d4944415478da63f8cfc0f01f000501020091d10a2c0000000049454e44ae426082"
        )
        img = XLImage(io.BytesIO(png))
        ws.add_image(img, f"A{10 + i}")
    wb.save(path)


def _ops(*specs) -> list:
    """快速构造 ExcelOperation 列表"""
    return [ExcelOperation(**spec) for spec in specs]


# ---------------------------------------------------------------------------
# set_cell
# ---------------------------------------------------------------------------
class TestSetCell:
    def test_set_cell_auto_number(self, file_service, temp_dir):
        """auto 类型：纯数字字符串应转为数字"""
        path = os.path.join(temp_dir, "t.xlsx")
        _make_xlsx(path)

        result = file_service.save_excel_edit("t.xlsx", _ops(
            {"op": "set_cell", "sheet": "Sheet1", "row": 2, "col": 2, "value": "30", "kind": "auto"},
        ))

        assert result.success is True
        ws = load_workbook(path)["Sheet1"]
        assert ws["B2"].value == 30

    def test_set_cell_auto_text(self, file_service, temp_dir):
        """auto 类型：普通文本保持字符串"""
        path = os.path.join(temp_dir, "t.xlsx")
        _make_xlsx(path)

        result = file_service.save_excel_edit("t.xlsx", _ops(
            {"op": "set_cell", "sheet": "Sheet1", "row": 2, "col": 1, "value": "Bob", "kind": "auto"},
        ))

        assert result.success is True
        assert load_workbook(path)["Sheet1"]["A2"].value == "Bob"

    def test_set_cell_str_keeps_leading_zero(self, file_service, temp_dir):
        """str 类型：前导零编号不应被转为数字"""
        path = os.path.join(temp_dir, "t.xlsx")
        _make_xlsx(path)

        result = file_service.save_excel_edit("t.xlsx", _ops(
            {"op": "set_cell", "sheet": "Sheet1", "row": 2, "col": 1, "value": "0012", "kind": "str"},
        ))

        assert result.success is True
        assert load_workbook(path)["Sheet1"]["A2"].value == "0012"

    def test_set_cell_bool(self, file_service, temp_dir):
        """bool 类型：TRUE/FALSE 转为布尔"""
        path = os.path.join(temp_dir, "t.xlsx")
        _make_xlsx(path)

        result = file_service.save_excel_edit("t.xlsx", _ops(
            {"op": "set_cell", "sheet": "Sheet1", "row": 1, "col": 3, "value": "TRUE", "kind": "bool"},
            {"op": "set_cell", "sheet": "Sheet1", "row": 2, "col": 3, "value": "false", "kind": "bool"},
        ))

        assert result.success is True
        ws = load_workbook(path)["Sheet1"]
        assert ws["C1"].value is True
        assert ws["C2"].value is False

    def test_set_cell_formula(self, file_service, temp_dir):
        """formula 类型：=开头字符串按公式写入"""
        path = os.path.join(temp_dir, "t.xlsx")
        _make_xlsx(path)

        result = file_service.save_excel_edit("t.xlsx", _ops(
            {"op": "set_cell", "sheet": "Sheet1", "row": 3, "col": 2, "value": "=SUM(B1:B2)", "kind": "formula"},
        ))

        assert result.success is True
        ws = load_workbook(path)["Sheet1"]
        assert ws["B3"].value == "=SUM(B1:B2)"

    def test_set_cell_nonexistent_sheet(self, file_service, temp_dir):
        """写入不存在的 sheet 应报错"""
        path = os.path.join(temp_dir, "t.xlsx")
        _make_xlsx(path)

        with pytest.raises(ValueError, match="工作表不存在"):
            file_service.save_excel_edit("t.xlsx", _ops(
                {"op": "set_cell", "sheet": "Nope", "row": 1, "col": 1, "value": "x", "kind": "auto"},
            ))


# ---------------------------------------------------------------------------
# 行列插入删除
# ---------------------------------------------------------------------------
class TestRowColOperations:
    def test_insert_row_shifts_existing(self, file_service, temp_dir):
        """插入行后原行应下移，并可在新行写入"""
        path = os.path.join(temp_dir, "t.xlsx")
        _make_xlsx(path)

        result = file_service.save_excel_edit("t.xlsx", _ops(
            {"op": "insert_row", "sheet": "Sheet1", "index": 2, "count": 1},
            {"op": "set_cell", "sheet": "Sheet1", "row": 2, "col": 1, "value": "新行", "kind": "str"},
        ))

        assert result.success is True
        ws = load_workbook(path)["Sheet1"]
        assert ws["A2"].value == "新行"
        assert ws["A3"].value == "Alice"

    def test_delete_row(self, file_service, temp_dir):
        """删除行后下方数据上移"""
        path = os.path.join(temp_dir, "t.xlsx")
        _make_xlsx(path)

        result = file_service.save_excel_edit("t.xlsx", _ops(
            {"op": "delete_row", "sheet": "Sheet1", "index": 1, "count": 1},
        ))

        assert result.success is True
        ws = load_workbook(path)["Sheet1"]
        assert ws["A1"].value == "Alice"

    def test_insert_col(self, file_service, temp_dir):
        """插入列后原列右移"""
        path = os.path.join(temp_dir, "t.xlsx")
        _make_xlsx(path)

        result = file_service.save_excel_edit("t.xlsx", _ops(
            {"op": "insert_col", "sheet": "Sheet1", "index": 2, "count": 1},
            {"op": "set_cell", "sheet": "Sheet1", "row": 1, "col": 2, "value": "中间列", "kind": "str"},
        ))

        assert result.success is True
        ws = load_workbook(path)["Sheet1"]
        assert ws["B1"].value == "中间列"
        assert ws["C1"].value == "Age"

    def test_delete_col(self, file_service, temp_dir):
        """删除列后右侧数据左移"""
        path = os.path.join(temp_dir, "t.xlsx")
        _make_xlsx(path)

        result = file_service.save_excel_edit("t.xlsx", _ops(
            {"op": "delete_col", "sheet": "Sheet1", "index": 1, "count": 1},
        ))

        assert result.success is True
        ws = load_workbook(path)["Sheet1"]
        assert ws["A1"].value == "Age"

    def test_ops_apply_in_order(self, file_service, temp_dir):
        """多条操作按顺序应用：插入后再按新坐标编辑"""
        path = os.path.join(temp_dir, "t.xlsx")
        _make_xlsx(path)

        result = file_service.save_excel_edit("t.xlsx", _ops(
            {"op": "insert_row", "sheet": "Sheet1", "index": 2, "count": 2},
            # 插入两行后，原第 2 行（Alice）变为第 4 行
            {"op": "set_cell", "sheet": "Sheet1", "row": 4, "col": 1, "value": "Alice2", "kind": "str"},
        ))

        assert result.success is True
        ws = load_workbook(path)["Sheet1"]
        assert ws["A2"].value is None
        assert ws["A4"].value == "Alice2"


# ---------------------------------------------------------------------------
# Sheet 操作
# ---------------------------------------------------------------------------
class TestSheetOperations:
    def test_add_rename_remove_sheet(self, file_service, temp_dir):
        """新增/重命名/删除 sheet 组合操作"""
        path = os.path.join(temp_dir, "t.xlsx")
        _make_xlsx(path)

        result = file_service.save_excel_edit("t.xlsx", _ops(
            {"op": "add_sheet", "title": "新增页"},
            {"op": "rename_sheet", "from_sheet": "新增页", "to_title": "改名页"},
        ))

        assert result.success is True
        assert load_workbook(path).sheetnames == ["Sheet1", "改名页"]

        result = file_service.save_excel_edit("t.xlsx", _ops(
            {"op": "remove_sheet", "sheet": "改名页"},
        ))
        assert result.success is True
        assert load_workbook(path).sheetnames == ["Sheet1"]

    def test_cannot_remove_last_sheet(self, file_service, temp_dir):
        """不允许删除最后一个 sheet"""
        path = os.path.join(temp_dir, "t.xlsx")
        _make_xlsx(path)

        with pytest.raises(ValueError, match="至少需要保留一个工作表"):
            file_service.save_excel_edit("t.xlsx", _ops(
                {"op": "remove_sheet", "sheet": "Sheet1"},
            ))

    def test_add_duplicate_sheet(self, file_service, temp_dir):
        """新增同名 sheet 应报错"""
        path = os.path.join(temp_dir, "t.xlsx")
        _make_xlsx(path)

        with pytest.raises(ValueError, match="工作表已存在"):
            file_service.save_excel_edit("t.xlsx", _ops(
                {"op": "add_sheet", "title": "Sheet1"},
            ))


# ---------------------------------------------------------------------------
# 图片保留（真实场景核心诉求）
# ---------------------------------------------------------------------------
class TestImagePreservation:
    def test_images_preserved_after_edit(self, file_service, temp_dir):
        """编辑保存后内嵌图片必须保留"""
        path = os.path.join(temp_dir, "with_img.xlsx")
        _make_xlsx(path, images=2)

        result = file_service.save_excel_edit("with_img.xlsx", _ops(
            {"op": "set_cell", "sheet": "Sheet1", "row": 2, "col": 1, "value": "改值", "kind": "str"},
            {"op": "insert_row", "sheet": "Sheet1", "index": 1, "count": 1},
        ))

        assert result.success is True
        wb = load_workbook(path)
        assert len(wb["Sheet1"]._images) == 2

    def test_no_tmp_file_left_behind(self, file_service, temp_dir):
        """保存成功后不应残留 .tmp 临时文件"""
        path = os.path.join(temp_dir, "t.xlsx")
        _make_xlsx(path)

        result = file_service.save_excel_edit("t.xlsx", _ops(
            {"op": "set_cell", "sheet": "Sheet1", "row": 1, "col": 1, "value": "x", "kind": "str"},
        ))

        assert result.success is True
        assert not os.path.exists(path + ".tmp")


# ---------------------------------------------------------------------------
# 输入防护
# ---------------------------------------------------------------------------
class TestInputValidation:
    def test_reject_non_xlsx(self, file_service, temp_dir):
        """非 .xlsx 文件应拒绝"""
        csv_path = os.path.join(temp_dir, "t.csv")
        with open(csv_path, "w", encoding="utf-8") as f:
            f.write("a,b\n1,2\n")

        result = file_service.save_excel_edit("t.csv", _ops(
            {"op": "set_cell", "sheet": "S", "row": 1, "col": 1, "value": "x", "kind": "str"},
        ))
        assert result.success is False
        assert ".xlsx" in result.message

    def test_reject_empty_operations(self, file_service, temp_dir):
        """空操作列表应拒绝"""
        path = os.path.join(temp_dir, "t.xlsx")
        _make_xlsx(path)

        result = file_service.save_excel_edit("t.xlsx", [])
        assert result.success is False

    def test_reject_missing_file(self, file_service, temp_dir):
        """文件不存在应返回失败"""
        result = file_service.save_excel_edit("ghost.xlsx", _ops(
            {"op": "set_cell", "sheet": "S", "row": 1, "col": 1, "value": "x", "kind": "str"},
        ))
        assert result.success is False
        assert "not found" in result.message.lower()

    def test_reject_path_escape(self, file_service, temp_dir):
        """路径逃逸（..）应被拒绝"""
        with pytest.raises(Exception):
            file_service.save_excel_edit("../outside.xlsx", _ops(
                {"op": "set_cell", "sheet": "S", "row": 1, "col": 1, "value": "x", "kind": "str"},
            ))

    def test_reject_missing_required_field(self, file_service, temp_dir):
        """set_cell 缺少 row 字段应报 ValueError"""
        path = os.path.join(temp_dir, "t.xlsx")
        _make_xlsx(path)

        with pytest.raises(ValueError, match="缺少必填字段"):
            file_service.save_excel_edit("t.xlsx", _ops(
                {"op": "set_cell", "sheet": "Sheet1", "col": 1, "value": "x", "kind": "str"},
            ))
